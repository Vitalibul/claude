import zipfile, re, copy, os, sys, json, subprocess, shutil
from lxml import etree
from docx import Document
from docxcompose.composer import Composer
from docx.enum.section import WD_SECTION

WN = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'; W = '{%s}' % WN
XS = '{http://www.w3.org/XML/1998/namespace}space'
def E(tag, **a):
    e = etree.Element(W + tag)
    for k, v in a.items(): e.set(W + k, str(v))
    return e
def sub(p, tag, **a):
    e = E(tag, **a); p.append(e); return e
def txt(e): return ''.join(t.text or '' for t in e.iter(W + 't'))

RPR_ORDER = ['rStyle', 'rFonts', 'b', 'bCs', 'i', 'iCs', 'caps', 'smallCaps', 'strike', 'dstrike', 'outline', 'shadow', 'emboss', 'imprint', 'noProof', 'snapToGrid', 'vanish', 'webHidden', 'color', 'spacing', 'w', 'kern', 'position', 'sz', 'szCs', 'highlight', 'u', 'effect', 'bdr', 'shd', 'fitText', 'vertAlign', 'rtl', 'cs', 'em', 'lang', 'eastAsianLayout', 'specVanish', 'oMath']
PPR_ORDER = ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl', 'numPr', 'suppressLineNumbers', 'pBdr', 'shd', 'tabs', 'suppressAutoHyphens', 'kinsoku', 'wordWrap', 'overflowPunct', 'topLinePunct', 'autoSpaceDE', 'autoSpaceDN', 'bidi', 'adjustRightInd', 'snapToGrid', 'spacing', 'ind', 'contextualSpacing', 'mirrorIndents', 'suppressOverlap', 'jc', 'textDirection', 'textAlignment', 'textboxTightWrap', 'outlineLvl', 'divId', 'cnfStyle', 'rPr', 'sectPr', 'pPrChange']
def put(parent, el, order):
    """insert el into parent respecting schema order; replaces same tag"""
    name = el.tag.split('}')[1]
    old = parent.find(W + name)
    if old is not None: parent.replace(old, el); return
    idx = order.index(name) if name in order else len(order)
    for i, ch in enumerate(parent):
        cn = ch.tag.split('}')[1]
        if cn in order and order.index(cn) > idx: ch.addprevious(el); return
    parent.append(el)
def get_pPr(p):
    pp = p.find(W + 'pPr')
    if pp is None: pp = E('pPr'); p.insert(0, pp)
    return pp
def get_rPr(r):
    rp = r.find(W + 'rPr')
    if rp is None: rp = E('rPr'); r.insert(0, rp)
    return rp

# ---------------------------------------------------------------- styles of a source doc
class Styles:
    def __init__(s, zf):
        s.x = etree.fromstring(zf.read('word/styles.xml'))
        s.by = {st.get(W + 'styleId'): st for st in s.x.iter(W + 'style')}
        dd = s.x.find(W + 'docDefaults/' + W + 'rPrDefault/' + W + 'rPr')
        s.defaults = {c.tag.split('}')[1]: c for c in dd} if dd is not None else {}
    def chain(s, sid, kind):
        out = {}; seen = 0
        while sid and sid in s.by and seen < 10:
            st = s.by[sid]; pr = st.find(W + kind)
            if pr is not None:
                for c in pr:
                    n = c.tag.split('}')[1]
                    if n not in out: out[n] = c
            b = st.find(W + 'basedOn'); sid = b.get(W + 'val') if b is not None else None; seen += 1
        return out

MASTER_DEF = {'font': 'Arial', 'sz': '20', 'color': None}

def bake_defaults(root, sty):
    """make runs independent of the source docDefaults (font / size / colour) that differ from the master"""
    dfl = sty.defaults
    dfont = dfl.get('rFonts'); dsz = dfl.get('sz'); dcol = dfl.get('color')
    need_font = dfont is not None and dfont.get(W + 'ascii') != MASTER_DEF['font']
    need_sz = dsz is not None and dsz.get(W + 'val') != MASTER_DEF['sz']
    need_col = dcol is not None
    if not (need_font or need_sz or need_col): return 0
    n = 0
    for p in root.iter(W + 'p'):
        ps = p.find(W + 'pPr/' + W + 'pStyle'); pst = sty.chain(ps.get(W + 'val'), 'rPr') if ps is not None else {}
        targets = list(p.iter(W + 'r')) + [None]
        for r in targets:
            if r is None:
                rp = get_pPr(p).find(W + 'rPr')
                if rp is None: rp = E('rPr'); put(get_pPr(p), rp, PPR_ORDER)
                have = {c.tag.split('}')[1] for c in rp}; cst = {}
            else:
                rp = r.find(W + 'rPr'); have = {c.tag.split('}')[1] for c in rp} if rp is not None else set()
                rs = rp.find(W + 'rStyle') if rp is not None else None
                cst = sty.chain(rs.get(W + 'val'), 'rPr') if rs is not None else {}
            def missing(tag): return tag not in have and tag not in pst and tag not in cst
            adds = []
            if need_font and missing('rFonts'): adds.append(copy.deepcopy(dfont))
            if need_sz and missing('sz'):
                adds.append(copy.deepcopy(dsz))
                if 'szCs' in dfl: adds.append(copy.deepcopy(dfl['szCs']))
            if need_col and missing('color'): adds.append(copy.deepcopy(dcol))
            if adds:
                if r is not None: rp = get_rPr(r)
                for a in adds: put(rp, a, RPR_ORDER)
                n += 1
    return n

def bake_style(p, sty, sid):
    """copy the formatting of paragraph style sid into p (direct formatting), drop the style"""
    pp = get_pPr(p)
    for tag, el in sty.chain(sid, 'pPr').items():
        if tag in ('outlineLvl', 'pStyle', 'rPr', 'numPr'): continue
        if pp.find(W + tag) is None: put(pp, copy.deepcopy(el), PPR_ORDER)
    rpr = sty.chain(sid, 'rPr')
    for r in list(p.iter(W + 'r')) + [None]:
        rp = (get_rPr(r) if r is not None else pp.find(W + 'rPr'))
        if rp is None: rp = E('rPr'); put(pp, rp, PPR_ORDER)
        for tag, el in rpr.items():
            if rp.find(W + tag) is None: put(rp, copy.deepcopy(el), RPR_ORDER)
    st = pp.find(W + 'pStyle')
    if st is not None: pp.remove(st)

def set_heading(p, level):
    pp = get_pPr(p); put(pp, E('pStyle', val=f'Heading{level}'), PPR_ORDER)

def in_table(p): return any(a.tag == W + 'tc' for a in p.iterancestors())

def prep(src, dst, variant_heading=True):
    zf = zipfile.ZipFile(src); sty = Styles(zf)
    parts = {}
    for n in zf.namelist():
        if n == 'word/document.xml' or re.match(r'word/(header|footer)\d*\.xml$', n): parts[n] = etree.fromstring(zf.read(n))
    stats = {'baked': 0}
    for n, root in parts.items(): stats['baked'] += bake_defaults(root, sty)
    doc = parts['word/document.xml']; body = doc.find(W + 'body')
    heads = []; mod_done = False
    for p in body.iter(W + 'p'):
        t = re.sub(r'\s+', ' ', txt(p)).strip()
        ps = p.find(W + 'pPr/' + W + 'pStyle'); sid = ps.get(W + 'val') if ps is not None else None
        if sid and (sid.startswith('Heading') or sid == 'Title'): bake_style(p, sty, sid)
        lvl = None
        if not mod_done and re.match(r'MODULE \d+ —', t): lvl = 1; mod_done = True
        elif re.match(r'SESSION \d+ —', t): lvl = 2
        elif re.match(r'Lesson \d+\.\d+\b.{0,20}—', t) and not in_table(p): lvl = 3
        elif re.match(r'TEST (CUMULATIV|FINAL)', t): lvl = 4
        elif variant_heading and re.match(r'VARIANTA: ', t): lvl = 4
        if lvl:
            set_heading(p, lvl); heads.append((lvl, t))
    zout = zipfile.ZipFile(dst + '.tmp', 'w', zipfile.ZIP_DEFLATED)
    for it in zf.infolist():
        data = zf.read(it.filename)
        if it.filename in parts: data = etree.tostring(parts[it.filename], xml_declaration=True, encoding='UTF-8', standalone=True)
        zout.writestr(it, data)
    zout.close(); os.replace(dst + '.tmp', dst)
    return heads, stats

# ---------------------------------------------------------------- sections
def src_page_props(path):
    x = etree.fromstring(zipfile.ZipFile(path).read('word/document.xml'))
    sp = x.find(W + 'body/' + W + 'sectPr')
    return {t: copy.deepcopy(sp.find(W + t)) for t in ('pgSz', 'pgMar', 'cols', 'docGrid') if sp.find(W + t) is not None}
def src_hf(path):
    d = Document(path); s = d.sections[0]
    return s.header._element, s.footer._element

SECT_ORDER = ['headerReference', 'footerReference', 'footnotePr', 'endnotePr', 'type', 'pgSz', 'pgMar', 'paperSrc', 'pgBorders', 'lnNumType', 'pgNumType', 'cols', 'formProt', 'vAlign', 'noEndnote', 'titlePg', 'textDirection', 'bidi', 'rtlGutter', 'docGrid', 'printerSettings', 'sectPrChange']
def apply_page_props(sp, props):
    for t, el in props.items(): put(sp, copy.deepcopy(el), SECT_ORDER)

def fill_hf(part_el, src_el, page_number):
    for ch in list(part_el): part_el.remove(ch)
    if src_el is not None:
        for ch in src_el:
            if ch.tag == W + 'p' and not txt(ch).strip() and page_number: continue
            part_el.append(copy.deepcopy(ch))
    if page_number:
        p = sub(part_el, 'p'); pp = sub(p, 'pPr'); sub(pp, 'spacing', before=40, after=0); sub(pp, 'jc', val='center')
        def r():
            rr = sub(p, 'r'); rp = sub(rr, 'rPr'); sub(rp, 'rFonts', ascii='Arial', hAnsi='Arial', cs='Arial'); sub(rp, 'color', val='555555'); sub(rp, 'sz', val=18); sub(rp, 'szCs', val=18); return rr
        sub(r(), 'fldChar', fldCharType='begin'); it = sub(r(), 'instrText'); it.text = ' PAGE '; it.set(XS, 'preserve')
        sub(r(), 'fldChar', fldCharType='separate'); t = sub(r(), 't'); t.text = '1'; sub(r(), 'fldChar', fldCharType='end')
    if len(part_el) == 0: sub(part_el, 'p')

# ---------------------------------------------------------------- TOC
def toc_styles(styles_root):
    for sid in ['Heading1', 'Heading2', 'Heading3', 'Heading4', 'TOC1', 'TOC2', 'TOC3', 'TOC4', 'TOCHeading']:
        old = styles_root.find(f"{W}style[@{W}styleId='{sid}']")
        if old is not None: styles_root.remove(old)
    def st(sid, name, ppr=None, rpr=None, extra=True):
        s = sub(styles_root, 'style', type='paragraph', styleId=sid); sub(s, 'name', val=name); sub(s, 'basedOn', val='Normal'); sub(s, 'next', val='Normal')
        sub(s, 'uiPriority', val=9 if sid.startswith('Heading') else 39)
        if sid.startswith('Heading'): sub(s, 'qFormat')
        else: sub(s, 'unhideWhenUsed')
        if ppr is not None: s.append(ppr)
        if rpr is not None: s.append(rpr)
    for i in range(1, 5):
        pp = E('pPr'); sub(pp, 'outlineLvl', val=i - 1); st(f'Heading{i}', f'heading {i}', pp)
    ind = {1: 0, 2: 240, 3: 480, 4: 240}
    for i in range(1, 5):
        pp = E('pPr'); tb = sub(pp, 'tabs'); sub(tb, 'tab', val='right', leader='dot', pos=9630)
        sub(pp, 'spacing', before=120 if i == 1 else 0, after=20)
        if ind[i]: sub(pp, 'ind', left=ind[i])
        rp = E('rPr')
        if i == 1: sub(rp, 'b'); sub(rp, 'color', val='1F5C45')
        if i == 4: sub(rp, 'i'); sub(rp, 'color', val='666666')
        sub(rp, 'sz', val=19 if i < 4 else 17); sub(rp, 'szCs', val=19 if i < 4 else 17)
        st(f'TOC{i}', f'toc {i}', pp, rp)

def toc_paragraphs(entries, nums):
    """entries: [(level, text, bm)] ; nums: bm -> displayed page (None for level 4)"""
    out = []
    for k, (lvl, text, bm) in enumerate(entries):
        p = E('p'); pp = sub(p, 'pPr'); sub(pp, 'pStyle', val=f'TOC{lvl}')
        if k == 0:
            r = sub(p, 'r'); sub(r, 'fldChar', fldCharType='begin')
            r = sub(p, 'r'); it = sub(r, 'instrText'); it.text = ' TOC \\o "1-4" \\h \\z \\u \\n "4-4" '; it.set(XS, 'preserve')
            r = sub(p, 'r'); sub(r, 'fldChar', fldCharType='separate')
        h = sub(p, 'hyperlink', anchor=bm, history=1)
        r = sub(h, 'r'); t = sub(r, 't'); t.text = text; t.set(XS, 'preserve')
        if lvl < 4:
            r = sub(h, 'r'); sub(r, 'tab')
            r = sub(h, 'r'); sub(r, 'fldChar', fldCharType='begin')
            r = sub(h, 'r'); it = sub(r, 'instrText'); it.text = f' PAGEREF {bm} \\h '; it.set(XS, 'preserve')
            r = sub(h, 'r'); sub(r, 'fldChar', fldCharType='separate')
            r = sub(h, 'r'); t = sub(r, 't'); t.text = str(nums.get(bm, 0))
            r = sub(h, 'r'); sub(r, 'fldChar', fldCharType='end')
        if k == len(entries) - 1:
            r = sub(p, 'r'); sub(r, 'fldChar', fldCharType='end')
        out.append(p)
    return out

# ---------------------------------------------------------------- build
def build(edition, modules, out_docx, work):
    os.makedirs(work, exist_ok=True)
    preps = []; allheads = []
    for m in modules:
        dst = os.path.join(work, os.path.basename(m))
        heads, stats = prep(m, dst); preps.append(dst); allheads.append(heads)
    master = Document(f'cover_{edition}.docx'); comp = Composer(master)
    for p in preps:
        comp.doc.add_section(WD_SECTION.NEW_PAGE); comp.append(Document(p))
    merged = os.path.join(work, 'merged.docx'); comp.save(merged)

    d = Document(merged); body = d.element.body
    # --- make the add_section paragraphs tiny
    sects = [p for p in body.iter(W + 'p') if p.find(W + 'pPr/' + W + 'sectPr') is not None]
    for p in sects:
        if not txt(p).strip():
            pp = p.find(W + 'pPr'); put(pp, E('spacing', before=0, after=0, line=20, lineRule='exact'), PPR_ORDER)
    # sectPr list in order: [cover, toc, mod1..modN-1] + body sectPr (modN)
    sps = [p.find(W + 'pPr/' + W + 'sectPr') for p in sects] + [body.find(W + 'sectPr')]
    assert len(sps) == 2 + len(modules), (len(sps), len(modules))
    kinds = ['cover', 'toc']
    # module ranges: elements between sect paragraphs
    top = list(body)
    def idx(e):
        while e.getparent() is not body: e = e.getparent()
        return top.index(e)
    bounds = [idx(p) for p in sects]  # index of paragraph that closes section i
    new_sps = []
    for k in range(len(modules)):
        mod_sp = sps[2 + k]; props = src_page_props(modules[k]); apply_page_props(mod_sp, props)
        lo = bounds[1 + k] + 1; hi = bounds[2 + k] if 2 + k < len(bounds) else len(top) - 1
        det = None
        for e in top[lo:hi]:
            if 'SECȚIUNE DETAȘABILĂ' in txt(e)[:60]: det = e; break
        assert det is not None, modules[k]
        # new section break before the detachable part
        bp = E('p'); pp = sub(bp, 'pPr'); put(pp, E('spacing', before=0, after=0, line=20, lineRule='exact'), PPR_ORDER)
        nsp = copy.deepcopy(mod_sp)
        for ch in list(nsp):
            if ch.tag in (W + 'headerReference', W + 'footerReference', W + 'pgNumType', W + 'titlePg'): nsp.remove(ch)
        put(nsp, E('type', val='nextPage'), SECT_ORDER)
        put(pp, nsp, PPR_ORDER); det.addprevious(bp)
        # a page break before the detachable heading is now redundant
        if det.tag == W + 'p':
            pb = det.find(W + 'pPr/' + W + 'pageBreakBefore')
            if pb is not None: pb.getparent().remove(pb)
        kinds += ['body', 'detach']
    for sp in [s for s in body.iter(W + 'sectPr')]:
        for ch in list(sp):
            if ch.tag in (W + 'headerReference', W + 'footerReference', W + 'titlePg'): sp.remove(ch)
    d.save(merged); d = Document(merged); body = d.element.body
    secs = d.sections; assert len(secs) == len(kinds), (len(secs), len(kinds))
    hfsrc = [src_hf(m) for m in preps]
    mi = -1
    for i, (s, kind) in enumerate(zip(secs, kinds)):
        s.header.is_linked_to_previous = False; s.footer.is_linked_to_previous = False
        s.different_first_page_header_footer = False
        if kind == 'body': mi += 1
        if kind in ('cover', 'toc'):
            fill_hf(s.header._element, None, False); fill_hf(s.footer._element, None, kind == 'toc')
        else:
            hsrc, fsrc = hfsrc[mi]
            fill_hf(s.header._element, hsrc, False); fill_hf(s.footer._element, fsrc, kind == 'body')
    # --- headings: bookmarks + TOC placeholder
    entries = []; bid = 900000
    for p in body.iter(W + 'p'):
        st = p.find(W + 'pPr/' + W + 'pStyle')
        if st is None or not re.match(r'Heading[1-4]$', st.get(W + 'val')): continue
        lvl = int(st.get(W + 'val')[-1]); t = re.sub(r'\s+', ' ', txt(p)).strip()
        bm = f'_Toc{bid}'; bs = E('bookmarkStart', id=bid, name=bm); be = E('bookmarkEnd', id=bid)
        pp = p.find(W + 'pPr'); pp.addnext(bs); p.append(be); bid += 1
        entries.append((lvl, t, bm))
    toc_styles(d.styles.element)
    ph = [p for p in body.iter(W + 'p') if txt(p).strip() == '@@TOC@@'][0]
    d.save(merged)
    return merged, entries, kinds, ph is not None

def put_toc(path, entries, nums, starts, out):
    d = Document(path); body = d.element.body
    ph = [p for p in body.iter(W + 'p') if txt(p).strip() == '@@TOC@@']
    if ph:
        ph = ph[0]
    else:
        ph = None
    old = [p for p in body.iter(W + 'p') if (p.find(W + 'pPr/' + W + 'pStyle') is not None and re.match(r'TOC[1-4]$', p.find(W + 'pPr/' + W + 'pStyle').get(W + 'val')))]
    anchor = ph if ph is not None else old[0]
    for p in toc_paragraphs(entries, nums): anchor.addprevious(p)
    if ph is not None: ph.getparent().remove(ph)
    for p in old: p.getparent().remove(p)
    # page number starts
    for sec, start in zip(d.sections, starts):
        sp = sec._sectPr
        pg = sp.find(W + 'pgNumType')
        if pg is not None: sp.remove(pg)
        if start is not None: put(sp, E('pgNumType', start=start), SECT_ORDER)
    d.save(out)
