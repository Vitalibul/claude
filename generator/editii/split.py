"""Split an edited edition into: the student book (no detachable sections) and the teacher book (only the detachable sections).
Text and pictures are not touched; only whole sections are removed."""
import sys, re, copy, subprocess, os
from lxml import etree
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import pymupdf

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
XS = '{http://www.w3.org/XML/1998/namespace}space'
def txt(e): return ''.join(t.text or '' for t in e.iter(W + 't'))
norm = lambda s: re.sub(r'\s+', ' ', s).strip()

def sections(body):
    secs = [[]]
    for e in body:
        if e.tag == W + 'sectPr': continue
        secs[-1].append(e)
        if e.tag == W + 'p' and e.find(W + 'pPr/' + W + 'sectPr') is not None: secs.append([])
    return secs
def is_detach(sec): return any('SECȚIUNE DETAȘABILĂ' in txt(e) for e in sec[:3])

def prune_media(doc):
    part = doc.part
    used = set()
    for el in part._element.iter():
        for k, v in el.attrib.items():
            if k.startswith(R): used.add(v)
    for rId, rel in list(part.rels.items()):
        if rel.reltype == RT.IMAGE and not rel.is_external and rId not in used:
            del part.rels[rId]

def remove_sections(doc, keep):
    body = doc.element.body; secs = sections(body)
    last_kept = max(i for i in range(len(secs)) if keep(i, secs[i]))
    for i, s in enumerate(secs):
        if keep(i, s): continue
        for e in s: body.remove(e)
    # if the trailing section (closed by the body sectPr) was removed, the last kept section becomes the final one
    if not keep(len(secs) - 1, secs[-1]):
        closing = secs[last_kept][-1]
        sp = closing.find(W + 'pPr/' + W + 'sectPr')
        old = body.find(W + 'sectPr'); body.replace(old, copy.deepcopy(sp))
        closing.find(W + 'pPr').remove(sp)
        if not txt(closing).strip() and closing.find('.//' + W + 'drawing') is None: body.remove(closing)
    return secs

def student(src, out):
    doc = Document(src); body = doc.element.body
    remove_sections(doc, lambda i, s: not is_detach(s))
    # page numbers: continuous (no fixed starts after the TOC section)
    sps = list(body.iter(W + 'sectPr'))
    first = True
    for sp in sps:
        pg = sp.find(W + 'pgNumType')
        if pg is not None and pg.get(W + 'start'):
            if first: first = False; continue
            sp.remove(pg)
    # TOC: drop entries whose target heading is gone (the tests)
    bms = {b.get(W + 'name') for b in body.iter(W + 'bookmarkStart')}
    toc = [p for p in body.iter(W + 'p') if p.find(W + 'pPr/' + W + 'pStyle') is not None and re.match(r'TOC\d$', p.find(W + 'pPr/' + W + 'pStyle').get(W + 'val'))]
    endrun = None
    for r in toc[-1].findall(W + 'r'):
        if r.find(W + 'fldChar') is not None and r.find(W + 'fldChar').get(W + 'fldCharType') == 'end': endrun = r
    kept = []
    for p in toc:
        h = p.find(W + 'hyperlink')
        if h is not None and h.get(W + 'anchor') not in bms:
            if p is toc[0]: raise Exception('first TOC entry removed')
            p.getparent().remove(p)
        else: kept.append(p)
    if endrun is not None and endrun.getparent() not in kept:
        kept[-1].append(endrun)
    prune_media(doc); doc.save(out)
    return kept

def teacher(src, out, title):
    doc = Document(src); body = doc.element.body
    secs = remove_sections(doc, lambda i, s: is_detach(s))
    # a title page in its own section, without header
    first = next(e for e in body if e.tag != W + 'sectPr')
    sp = copy.deepcopy(body.find(W + 'sectPr'))
    for ch in list(sp):
        if ch.tag in (W + 'headerReference', W + 'footerReference', W + 'pgNumType'): sp.remove(ch)
    def para(text, sz, bold=True, color='1F5C45', after=200):
        p = etree.Element(W + 'p'); pp = etree.SubElement(p, W + 'pPr')
        etree.SubElement(pp, W + 'spacing').set(W + 'after', str(after)); etree.SubElement(pp, W + 'jc').set(W + 'val', 'center')
        r = etree.SubElement(p, W + 'r'); rp = etree.SubElement(r, W + 'rPr')
        f = etree.SubElement(rp, W + 'rFonts'); [f.set(W + a, 'Arial') for a in ('ascii', 'hAnsi', 'cs')]
        if bold: etree.SubElement(rp, W + 'b')
        etree.SubElement(rp, W + 'color').set(W + 'val', color); etree.SubElement(rp, W + 'sz').set(W + 'val', str(sz))
        t = etree.SubElement(r, W + 't'); t.text = text; t.set(XS, 'preserve'); return p
    items = [para('', 20, after=2400), para('LIMBA ROMÂNĂ PENTRU STRĂINI', 44, after=160), para('EXEMPLAR FORMATOR', 36, after=160),
             para('Testele cumulative, fișele de evidență și cheile de răspunsuri ale Modulelor 1–14', 24, False, '333333', 120),
             para(title, 24, False, '555555', 120)]
    last = para('', 20, after=0); last.find(W + 'pPr').append(sp); items.append(last)
    for e in items: first.addprevious(e)
    # no page numbers in this book: nothing to fix; prune unused pictures
    prune_media(doc); doc.save(out)

def render(docx):
    out = os.path.dirname(os.path.abspath(docx))
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', out, docx], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=1500)
    return pymupdf.open(docx[:-5] + '.pdf')

def fix_toc_numbers(path):
    """recompute the cached page numbers of the TOC from the LibreOffice layout (continuous numbering from the TOC page)"""
    pdf = render(path); pages = [norm(p.get_text()) for p in pdf]
    doc = Document(path); body = doc.element.body
    heads = []
    for p in body.iter(W + 'p'):
        st = p.find(W + 'pPr/' + W + 'pStyle')
        if st is not None and re.match(r'Heading[1-3]$', st.get(W + 'val')):
            b = p.find(W + 'bookmarkStart')
            if b is not None: heads.append((b.get(W + 'name'), norm(txt(p))))
    toc0 = next(i for i, t in enumerate(pages) if 'CUPRINS DETALIAT' in t)
    i = toc0 + 1
    while not ('SUPORT DE CURS' in pages[i] and 'LIMBA ROMÂNĂ PENTRU STRĂINI' in pages[i]): i += 1
    cur = i; num = {}
    for bm, t in heads:
        key = t[:28]; j = cur
        while j < len(pages) and key not in pages[j]: j += 1
        if j >= len(pages): print('not found', t); continue
        num[bm] = j - toc0 + 1; cur = j
    n = 0
    for h in body.iter(W + 'hyperlink'):
        bm = h.get(W + 'anchor')
        if bm not in num: continue
        runs = h.findall(W + 'r'); state = None
        for r in runs:
            fc = r.find(W + 'fldChar')
            if fc is not None: state = fc.get(W + 'fldCharType'); continue
            if state == 'separate' and r.find(W + 't') is not None:
                r.find(W + 't').text = str(num[bm]); n += 1
    doc.save(path)
    return len(pdf), n

if __name__ == '__main__':
    src, stud, teach, title = sys.argv[1:5]
    student(src, stud)
    pages, n = fix_toc_numbers(stud)
    fix_toc_numbers(stud)  # second pass: TOC length may shift pages
    teacher(src, teach, title)
    print('student pages', pages, 'toc numbers', n)
