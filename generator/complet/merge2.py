from docx import Document
from docxcompose.composer import Composer
from docx.enum.section import WD_SECTION
import glob, copy
R='/home/user/claude/'
order=['Modul_1_','Modul_2_','Modul_3_','Modul_4_','Modul_5_','Modul_6_','Modul_7_','Modul_8_',
 'Modul_9_Domeniul_de_activitate_Constructii','Modul_9_Domeniul_de_activitate_HoReCa','Modul_9_Domeniul_de_activitate_Transport',
 'Modul_9_Domeniul_de_activitate_Comert','Modul_9_Domeniul_de_activitate_Productie','Modul_10_','Modul_11_','Modul_12_','Modul_13_','Modul_14_']
files=[]
for p in order:
    m=glob.glob(R+p+'*.docx'); assert len(m)==1,(p,m); files.append(m[0])
master=Document('cover.docx'); comp=Composer(master)
for f in files:
    comp.doc.add_section(WD_SECTION.NEW_PAGE)
    comp.append(Document(f))
comp.save('full.docx')
# copy each module's original header/footer XML verbatim into its section
d=Document('full.docx')
def copy_part(dst, src):
    dst.is_linked_to_previous=False
    de=dst._element
    for ch in list(de): de.remove(ch)
    for ch in src._element: de.append(copy.deepcopy(ch))
for i,f in enumerate(files, start=1):
    s=Document(f).sections[0]
    copy_part(d.sections[i].header, s.header); copy_part(d.sections[i].footer, s.footer)
d.save('full.docx'); print('ok', len(d.sections))
